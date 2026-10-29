import io
import re


class ResumeParseError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as e:
            raise ResumeParseError(f"Could not read PDF: {e}")
    elif name.endswith(".docx"):
        import docx

        try:
            doc = docx.Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        except Exception as e:
            raise ResumeParseError(f"Could not read DOCX: {e}")
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ResumeParseError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    if len(text) < 100:
        raise ResumeParseError("Very little text found. If this is a scanned PDF, export a text-based PDF or DOCX.")
    return text
